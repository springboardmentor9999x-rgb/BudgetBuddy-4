from io import BytesIO
from datetime import datetime

from openpyxl import load_workbook
from app.database import SessionLocal
from app.models.expense import Expense
from app.models.income import Income


def test_empty_monthly_report_and_exports_are_valid(client, headers_a):
    report = client.get("/reports/monthly", headers=headers_a)
    assert report.status_code == 200
    assert report.json()["summary"]["total_income"] == 0.0
    assert report.json()["transactions"] == []

    assert client.get("/reports/export/pdf", headers=headers_a).status_code == 200

    from datetime import datetime, timedelta, timezone
    from app.database import SessionLocal
    from app.models.subscription import Subscription
    db = SessionLocal()
    user_id = client.get("/auth/me", headers=headers_a).json()["id"]
    db.add(Subscription(user_id=user_id, plan="premium_monthly", status="active", starts_at=datetime.now(timezone.utc), expires_at=datetime.now(timezone.utc) + timedelta(days=30)))
    db.commit()
    db.close()

    pdf = client.get("/reports/export/pdf", headers=headers_a)
    assert pdf.status_code == 200
    assert pdf.headers["content-type"].startswith("application/pdf")
    assert pdf.content.startswith(b"%PDF")

    excel = client.get("/reports/export/excel", headers=headers_a)
    assert excel.status_code == 200
    assert excel.headers["content-type"].startswith("application/vnd.openxmlformats")
    workbook = load_workbook(BytesIO(excel.content), read_only=True)
    assert workbook.sheetnames == ["Summary", "Transactions"]
    assert "Reference" not in [cell.value for cell in next(workbook["Transactions"].iter_rows())]
    assert client.get("/reports/export/pdf", headers=headers_a).status_code == 200
    assert client.get("/reports/monthly", headers=headers_a).json()["export_allowance"]["limit"] is None


def test_monthly_report_contains_only_current_users_transactions(client, headers_a, headers_b):
    income = client.post(
        "/income/",
        headers=headers_a,
        json={
            "source": "Salary",
            "amount": 3000,
            "description": "Monthly pay",
            "bank_account": "Test Bank 1234",
        },
    )
    expense = client.post(
        "/expenses/",
        headers=headers_a,
        json={"category": "Food", "amount": 450, "description": "Groceries", "bank_account": "Test Bank 1234"},
    )
    assert income.status_code == 200
    assert expense.status_code == 200

    report = client.get("/reports/monthly", headers=headers_a).json()
    assert report["summary"]["total_income"] == 3000.0
    assert report["summary"]["total_expenses"] == 450.0
    assert report["spending_by_category"] == [{"category": "Food", "total": 450.0}]
    assert len(report["transactions"]) == 2

    other_user_report = client.get("/reports/monthly", headers=headers_b).json()
    assert other_user_report["transactions"] == []


def test_statement_uses_prior_cash_flow_as_opening_balance_and_ignores_legacy_negative_income(client, headers_a, user_a):
    now = datetime.now()
    previous_year = now.year - (now.month == 1)
    previous_month = 12 if now.month == 1 else now.month - 1
    db = SessionLocal()
    try:
        db.add_all([
            Income(user_id=user_a.id, source="Previous salary", amount=1000, description="Prior income", date=datetime(previous_year, previous_month, 10)),
            Expense(user_id=user_a.id, category="Food", amount=200, description="Prior expense", date=datetime(previous_year, previous_month, 11)),
            Income(user_id=user_a.id, source="Savings contribution — legacy", amount=-300, description="Legacy transfer", date=datetime(now.year, now.month, 1)),
        ])
        db.commit()
    finally:
        db.close()

    report = client.get(f"/reports/monthly?month={now.month}&year={now.year}", headers=headers_a).json()
    assert report["summary"]["opening_balance"] == 800.0
    assert report["summary"]["total_income"] == 0.0
    assert report["summary"]["closing_balance"] == 800.0
    assert report["transactions"] == []


def test_free_exports_share_lifetime_limit(client, headers_a, headers_b):
    assert client.get("/reports/monthly", headers=headers_a).json()["export_allowance"] == {"limit": 2, "remaining": 2}
    assert client.get("/reports/export/pdf", headers=headers_a).status_code == 200
    assert client.get("/reports/export/excel?month=1&year=2025", headers=headers_a).status_code == 200
    for format in ("pdf", "excel"):
        response = client.get(f"/reports/export/{format}", headers=headers_a)
        assert response.status_code == 403
        assert "2 free report exports" in response.json()["detail"]
    assert client.get("/reports/monthly", headers=headers_a).json()["export_allowance"]["remaining"] == 0
    assert client.get("/reports/export/pdf", headers=headers_b).status_code == 200


def test_invalid_export_does_not_consume_allowance(client, headers_a):
    assert client.get("/reports/export/pdf?month=13", headers=headers_a).status_code == 422
    assert client.get("/reports/monthly", headers=headers_a).json()["export_allowance"]["remaining"] == 2


def test_admin_exports_are_unlimited(client, headers_a, user_a):
    from app.models.user import User
    with SessionLocal() as db:
        db.query(User).filter(User.id == user_a.id).update({"role": "admin"})
        db.commit()
    for _ in range(3):
        assert client.get("/reports/export/pdf", headers=headers_a).status_code == 200
    assert client.get("/reports/monthly", headers=headers_a).json()["export_allowance"]["limit"] is None


def test_failed_generation_does_not_consume_allowance(client, headers_a, monkeypatch):
    import pytest
    from app.routers import reports

    def fail_build(*args, **kwargs):
        raise RuntimeError("Generation failed")

    monkeypatch.setattr(reports.SimpleDocTemplate, "build", fail_build)
    with pytest.raises(RuntimeError, match="Generation failed"):
        client.get("/reports/export/pdf", headers=headers_a)
    assert client.get("/reports/monthly", headers=headers_a).json()["export_allowance"]["remaining"] == 2
